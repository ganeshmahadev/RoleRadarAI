from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.services.company_normalization import build_eures_search_url, normalize_company_name

SEED = Path(__file__).resolve().parents[3] / "data" / "siri_certified_companies_eures_queue.xlsx"


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("3Shape A/S", "3shape"),
        ("&TRADITION A/S", "&tradition"),
        ("A.P. MØLLER HOLDING A/S", "a p møller holding"),
        ("Aalborg Håndbold A/S", "aalborg handbold"),
        ("BLÜCHER METAL A/S", "blucher metal"),
        ("A+ Siam Sushi Restaurant ApS", "a+ siam sushi restaurant"),
        ("A/S Bryggeriet Vestfyen", "a s bryggeriet vestfyen"),
        ("Arla Foods amba", "arla foods"),
        ("BLUE FOX - HERNING A/S", "blue fox herning"),
    ],
)
def test_normalize_company_name(name: str, expected: str) -> None:
    assert normalize_company_name(name) == expected


def test_normalization_reproduces_seed_workbook() -> None:
    rows = list(load_workbook(SEED, read_only=True)["SIRI Companies"].iter_rows(values_only=True))
    mismatches = [(r[1], r[4]) for r in rows[1:] if normalize_company_name(str(r[1])) != r[4]]
    assert len(rows) - 1 == 982
    assert mismatches == []


def test_eures_url_reproduces_seed_workbook() -> None:
    rows = list(load_workbook(SEED, read_only=True)["EURES Queue"].iter_rows(values_only=True))
    mismatches = [r[1] for r in rows[1:] if build_eures_search_url(str(r[1])) != r[8]]
    assert len(rows) - 1 == 982
    assert mismatches == []


def test_eures_url_encodes_special_characters() -> None:
    url = build_eures_search_url("&TRADITION A/S")
    assert "keywordsEverywhere=%26TRADITION+A%2FS&" in url
    assert url.startswith("https://europa.eu/eures/portal/jv-se/search?")
