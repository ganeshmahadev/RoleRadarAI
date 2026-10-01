"""Anonymized sample resumes generated in-process (never real personal data)."""

import io

import pymupdf
from docx import Document

SAMPLE_LINES = [
    "Alex Example",
    "Machine Learning Engineer - Aarhus, Denmark",
    "Experience: 5 years Python, PyTorch, FastAPI; built RAG systems.",
    "Languages: English (C2), Danish (A2)",
]


def make_pdf(lines: list[str] = SAMPLE_LINES, *, pages: int = 1) -> bytes:
    doc = pymupdf.open()
    for page_no in range(pages):
        page = doc.new_page()
        y = 72
        for line in lines:
            page.insert_text((72, y), f"{line}" if pages == 1 else f"{line} (p{page_no + 1})")
            y += 18
    data: bytes = doc.tobytes()
    doc.close()
    return data


def make_blank_pdf() -> bytes:
    """A PDF with no text layer, like a scanned resume."""
    doc = pymupdf.open()
    doc.new_page()
    data: bytes = doc.tobytes()
    doc.close()
    return data


def make_locked_pdf() -> bytes:
    """A password-protected PDF."""
    doc = pymupdf.open(stream=make_pdf(), filetype="pdf")
    data: bytes = doc.tobytes(
        encryption=getattr(pymupdf, "PDF_ENCRYPT_AES_256"),  # noqa: B009  (missing from stubs)
        user_pw="secret",
        owner_pw="owner",
    )
    doc.close()
    return data


def make_docx(lines: list[str] = SAMPLE_LINES, *, table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for line in lines:
        document.add_paragraph(line)
    if table:
        grid = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                grid.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
