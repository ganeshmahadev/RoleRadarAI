"""Import the SIRI Fast-track certified-company workbook (PRD §7.1).

Rules:
- CVR values are read as text and preserved exactly; numeric CVR cells are rejected.
- Import is idempotent: upsert by CVR.
- Re-import updates only SIRI-owned fields and never touches EURES workflow state.
"""

from dataclasses import dataclass
from datetime import UTC, date, datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import Company, EuresStatus
from app.services.company_normalization import build_eures_search_url, normalize_company_name

SIRI_SHEET = "SIRI Companies"
COLUMNS = {
    "position": "ID",
    "company_name": "Company Name (Official)",
    "cvr": "CVR No.",
    "siri_certified": "SIRI Certified",
    "siri_last_seen_at": "Source Last Updated",
    "siri_source_url": "Source URL",
}
_CERTIFIED_VALUES = {"yes": True, "no": False}


class SiriImportError(ValueError):
    """The workbook is invalid; nothing was imported."""

    def __init__(self, errors: list[str]) -> None:
        super().__init__("; ".join(errors[:10]) + (" …" if len(errors) > 10 else ""))
        self.errors = errors


@dataclass(frozen=True)
class SiriCompanyRow:
    position: int
    company_name: str
    cvr: str
    siri_certified: bool
    siri_last_seen_at: datetime | None
    siri_source_url: str | None


@dataclass(frozen=True)
class ImportResult:
    total: int
    created: int
    updated: int
    unchanged: int


def _as_utc(value: object) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=UTC)
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day, tzinfo=UTC)
    raise TypeError(f"expected a date, got {type(value).__name__}")


def read_siri_workbook(path: Path) -> list[SiriCompanyRow]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if SIRI_SHEET not in workbook.sheetnames:
        raise SiriImportError([f"sheet {SIRI_SHEET!r} not found"])
    rows = workbook[SIRI_SHEET].iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows, ())]
    missing = [name for name in COLUMNS.values() if name not in header]
    if missing:
        raise SiriImportError([f"missing column {name!r}" for name in missing])
    index = {field: header.index(name) for field, name in COLUMNS.items()}

    parsed: list[SiriCompanyRow] = []
    errors: list[str] = []
    seen_cvrs: dict[str, int] = {}
    for line, row in enumerate(rows, start=2):
        if all(value is None for value in row):
            continue
        cvr = row[index["cvr"]]
        name = row[index["company_name"]]
        position = row[index["position"]]
        certified = row[index["siri_certified"]]
        row_errors: list[str] = []
        if not isinstance(cvr, str) or not cvr:
            # Never coerce: a numeric cell may already have lost leading zeros.
            row_errors.append(f"row {line}: CVR must be a non-empty text cell, got {cvr!r}")
        elif cvr in seen_cvrs:
            row_errors.append(
                f"row {line}: duplicate CVR {cvr!r} (first seen row {seen_cvrs[cvr]})"
            )
        if not isinstance(name, str) or not name.strip():
            row_errors.append(f"row {line}: company name is empty")
        if not isinstance(position, int):
            row_errors.append(f"row {line}: ID must be an integer, got {position!r}")
        if not isinstance(certified, str) or certified.strip().lower() not in _CERTIFIED_VALUES:
            row_errors.append(f"row {line}: SIRI Certified must be Yes/No, got {certified!r}")
        try:
            last_seen = _as_utc(row[index["siri_last_seen_at"]])
        except TypeError as exc:
            row_errors.append(f"row {line}: Source Last Updated {exc}")
        if row_errors:
            errors.extend(row_errors)
            continue
        assert isinstance(cvr, str) and isinstance(name, str) and isinstance(position, int)
        assert isinstance(certified, str)
        seen_cvrs[cvr] = line
        source_url = row[index["siri_source_url"]]
        parsed.append(
            SiriCompanyRow(
                position=position,
                company_name=name,
                cvr=cvr,
                siri_certified=_CERTIFIED_VALUES[certified.strip().lower()],
                siri_last_seen_at=last_seen,
                siri_source_url=str(source_url) if source_url else None,
            )
        )
    if errors:
        raise SiriImportError(errors)
    return parsed


def _siri_fields(row: SiriCompanyRow) -> dict[str, object]:
    """Fields owned by the SIRI source. EURES workflow and user-entered fields are excluded."""
    return {
        "company_name": row.company_name,
        "normalized_name": normalize_company_name(row.company_name),
        "source_position": row.position,
        "siri_certified": row.siri_certified,
        "siri_last_seen_at": row.siri_last_seen_at,
        "siri_source_url": row.siri_source_url,
        "eures_search_url": build_eures_search_url(row.company_name),
    }


async def import_siri_companies(session: AsyncSession, rows: list[SiriCompanyRow]) -> ImportResult:
    existing = {
        company.cvr: company
        for company in (
            await session.execute(select(Company).where(Company.cvr.in_([r.cvr for r in rows])))
        ).scalars()
    }
    created = updated = unchanged = 0
    for row in rows:
        fields = _siri_fields(row)
        company = existing.get(row.cvr)
        if company is None:
            session.add(Company(cvr=row.cvr, eures_status=EuresStatus.NOT_CHECKED, **fields))
            created += 1
            continue
        changed = False
        for key, value in fields.items():
            if getattr(company, key) != value:
                setattr(company, key, value)
                changed = True
        if changed:
            updated += 1
        else:
            unchanged += 1
    await session.commit()
    return ImportResult(total=len(rows), created=created, updated=updated, unchanged=unchanged)
